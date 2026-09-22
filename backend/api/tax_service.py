import csv
import io
from datetime import date
from decimal import Decimal
from django.db.models import Q, Sum
from .models import Transaction, Category, TaxFilterPreset

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


SYSTEM_AEAT_PRESETS = [
    {
        'name': 'Rendimientos del Trabajo (Nóminas)',
        'aeat_box': 'Casillas 0001-0025: Rendimientos del trabajo',
        'filters': {
            'is_income': True,
            'categories': ['Nómina', 'Ingresos del trabajo', 'Sueldo'],
            'aeat_codes': ['TRABAJO_0001'],
        }
    },
    {
        'name': 'Donativos y Aportaciones a ONG (Ley 49/2002)',
        'aeat_box': 'Casillas 0722-0724: Donaciones a entidades sin fines de lucro',
        'filters': {
            'is_income': False,
            'categories': ['Donaciones', 'ONG', 'Ayuda humanitaria', 'Cuotas de beneficencia'],
            'aeat_codes': ['DONACION_0722'],
            'is_tax_deductible': True
        }
    },
    {
        'name': 'Inversión en Vivienda Habitual (Hipoteca pre-2013)',
        'aeat_box': 'Casillas 0698-0700: Deducción por adquisición de vivienda habitual',
        'filters': {
            'is_income': False,
            'categories': ['Hipoteca', 'Amortización hipoteca', 'Seguro vida hipoteca', 'Seguro hogar hipoteca'],
            'aeat_codes': ['HIPOTECA_0698'],
            'is_tax_deductible': True
        }
    },
    {
        'name': 'Arrendamiento de Vivienda Habitual (Inquilino)',
        'aeat_box': 'Casilla 0102 / Deducciones autonómicas por alquiler',
        'filters': {
            'is_income': False,
            'categories': ['Alquiler', 'Renta alquiler', 'Vivienda'],
            'aeat_codes': ['ALQUILER_0102'],
            'is_tax_deductible': True
        }
    },
    {
        'name': 'Aportaciones a Planes de Pensiones',
        'aeat_box': 'Casilla 0465: Reducciones por aportaciones a sistemas de previsión social',
        'filters': {
            'is_income': False,
            'categories': ['Plan de pensiones', 'Fondo de pensiones', 'Previsión social'],
            'aeat_codes': ['PENSION_0465'],
            'is_tax_deductible': True
        }
    },
    {
        'name': 'Rendimientos del Capital Mobiliario',
        'aeat_box': 'Casillas 0029-0036: Intereses de cuentas, depósitos y dividendos',
        'filters': {
            'is_income': True,
            'categories': ['Dividendos', 'Intereses bancarios', 'Rendimientos capital'],
            'aeat_codes': ['DIVIDENDOS_0029']
        }
    },
    {
        'name': 'Ganancias y Pérdidas Patrimoniales',
        'aeat_box': 'Casillas 0380-0410: Ganancias/pérdidas por transmisión de elementos',
        'filters': {
            'categories': ['Inversiones', 'Criptomonedas', 'Acciones', 'Fondos de inversión'],
            'aeat_codes': ['PATRIMONIAL_0380']
        }
    },
]


def ensure_system_presets(user):
    """Crea los presets oficiales de la AEAT para el usuario si no existen."""
    for preset_data in SYSTEM_AEAT_PRESETS:
        TaxFilterPreset.objects.get_or_create(
            user=user,
            name=preset_data['name'],
            defaults={
                'aeat_box': preset_data['aeat_box'],
                'filters': preset_data['filters'],
                'is_system_preset': True
            }
        )


def get_tax_summary(user, year: int) -> dict:
    """
    Calcula el balance y desglose de la Declaración de la Renta para un ejercicio fiscal.
    Aplica las normas de desgravación de la AEAT para particulares.
    """
    ensure_system_presets(user)

    # Transacciones pertenecientes al año fiscal (por tax_year explícito o fecha)
    txs = Transaction.objects.filter(
        user=user,
        is_deleted=False
    ).filter(
        Q(tax_year=year) | (Q(tax_year__isnull=True) & Q(date__year=year))
    ).select_related('category', 'account').prefetch_related('documents')

    total_transactions = txs.count()
    total_with_document = 0
    total_income = Decimal('0.00')
    total_expense = Decimal('0.00')

    # Desgloses específicos
    donations_base = Decimal('0.00')
    mortgage_base = Decimal('0.00')
    rent_base = Decimal('0.00')
    pension_base = Decimal('0.00')
    work_income = Decimal('0.00')
    capital_income = Decimal('0.00')

    # Diccionario para agrupar por categoría / AEAT
    categories_breakdown = {}

    for tx in txs:
        docs_count = tx.documents.count()
        if docs_count > 0:
            total_with_document += 1

        amt = abs(tx.amount)
        cat_name = tx.category.name if tx.category else 'Sin categoría'
        aeat_code = tx.category.aeat_code if tx.category and tx.category.aeat_code else ''
        is_deductible = tx.is_tax_deductible or (tx.category and tx.category.tax_deductible)

        if tx.type == 'income':
            total_income += amt
            # Comprobar rendimientos de trabajo
            if aeat_code == 'TRABAJO_0001' or 'nómina' in cat_name.lower() or 'trabajo' in cat_name.lower():
                work_income += amt
            elif aeat_code == 'DIVIDENDOS_0029' or 'dividendo' in cat_name.lower() or 'interes' in cat_name.lower():
                capital_income += amt
        else:
            total_expense += amt
            # Donaciones
            if aeat_code == 'DONACION_0722' or 'donaci' in cat_name.lower() or 'ong' in cat_name.lower():
                donations_base += amt
            # Hipoteca pre-2013
            elif aeat_code == 'HIPOTECA_0698' or 'hipoteca' in cat_name.lower():
                mortgage_base += amt
            # Alquiler vivienda
            elif aeat_code == 'ALQUILER_0102' or 'alquiler' in cat_name.lower():
                rent_base += amt
            # Planes de pensiones
            elif aeat_code == 'PENSION_0465' or 'pension' in cat_name.lower():
                pension_base += amt

        # Acumular por categoría
        if cat_name not in categories_breakdown:
            categories_breakdown[cat_name] = {
                'category_name': cat_name,
                'category_color': tx.category.color if tx.category else '#cccccc',
                'category_icon': tx.category.icon if tx.category else 'receipt',
                'aeat_code': aeat_code,
                'is_income': tx.type == 'income',
                'is_deductible': is_deductible,
                'total_amount': Decimal('0.00'),
                'count': 0,
                'with_document_count': 0
            }
        categories_breakdown[cat_name]['total_amount'] += amt
        categories_breakdown[cat_name]['count'] += 1
        if docs_count > 0:
            categories_breakdown[cat_name]['with_document_count'] += 1

    # Cálculo fiscal deducción donaciones (Ley 49/2002):
    # Primeros 250€ al 80%, exceso al 40%
    donation_deduction_tier1 = min(donations_base, Decimal('250.00')) * Decimal('0.80')
    donation_deduction_tier2 = max(Decimal('0.00'), donations_base - Decimal('250.00')) * Decimal('0.40')
    donations_estimated_deduction = donation_deduction_tier1 + donation_deduction_tier2

    # Cálculo fiscal hipoteca pre-2013: 15% hasta base máxima de 9.040€ (max deducción 1.356€)
    mortgage_capped_base = min(mortgage_base, Decimal('9040.00'))
    mortgage_estimated_deduction = mortgage_capped_base * Decimal('0.15')

    # Planes de pensiones: límite reducción 1.500€
    pension_capped_base = min(pension_base, Decimal('1500.00'))

    # Cobertura documental
    coverage_pct = round((total_with_document / total_transactions * 100), 1) if total_transactions > 0 else 0.0

    return {
        'tax_year': year,
        'metrics': {
            'total_income': float(total_income),
            'total_expense': float(total_expense),
            'work_income': float(work_income),
            'capital_income': float(capital_income),
            'donations_base': float(donations_base),
            'donations_deduction': float(donations_estimated_deduction),
            'mortgage_base': float(mortgage_base),
            'mortgage_deduction': float(mortgage_estimated_deduction),
            'rent_base': float(rent_base),
            'pension_base': float(pension_base),
            'pension_capped_base': float(pension_capped_base),
            'total_transactions': total_transactions,
            'total_with_document': total_with_document,
            'document_coverage_percentage': coverage_pct,
        },
        'categories_breakdown': sorted(list(categories_breakdown.values()), key=lambda x: x['total_amount'], reverse=True)
    }


def export_tax_excel(user, year: int) -> io.BytesIO:
    """Genera informe Excel (.xlsx) estructurado para la Declaración de la Renta."""
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl no está disponible para exportar a Excel.")

    summary = get_tax_summary(user, year)
    metrics = summary['metrics']

    wb = openpyxl.Workbook()
    
    # ----------------------------------------------------
    # HOJA 1: Resumen IRPF por Casillas AEAT
    # ----------------------------------------------------
    ws_summary = wb.active
    ws_summary.title = f"Resumen Renta {year}"
    ws_summary.views.sheetView[0].showGridLines = True

    # Estilos
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    section_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    bold_font = Font(name="Calibri", size=10, bold=True)
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Título
    ws_summary['A1'] = f"INFORME FISCAL PARA DECLARACIÓN DE LA RENTA — EJERCICIO {year}"
    ws_summary['A1'].font = title_font
    ws_summary['A2'] = f"Contribuyente: {user.username} ({user.email}) | Generado por FinancIAs el {date.today().strftime('%d/%m/%Y')}"
    ws_summary['A2'].font = Font(name="Calibri", size=10, italic=True, color="64748B")

    # Tabla Resumen Casillas AEAT
    ws_summary['A4'] = "Concepto / Casilla AEAT"
    ws_summary['B4'] = "Base Computable (€)"
    ws_summary['C4'] = "Deducción Estimada (€)"
    ws_summary['D4'] = "Normativa y Límite Aplicable"

    for col in ['A4', 'B4', 'C4', 'D4']:
        ws_summary[col].fill = header_fill
        ws_summary[col].font = header_font
        ws_summary[col].alignment = Alignment(horizontal="left" if col in ('A4', 'D4') else "right", vertical="center")

    rows_data = [
        ("Casillas 0001-0025: Rendimientos del Trabajo (Nóminas)", metrics['work_income'], "-", "Base imponible general"),
        ("Casillas 0722-0724: Donativos y Aportaciones a ONG", metrics['donations_base'], metrics['donations_deduction'], "80% primeros 250€, 40% exceso (Ley 49/2002)"),
        ("Casillas 0698-0700: Deducción Vivienda Habitual (Pre-2013)", metrics['mortgage_base'], metrics['mortgage_deduction'], "15% sobre base máxima 9.040€ (Máx 1.356€)"),
        ("Casilla 0102: Arrendamiento Vivienda Habitual", metrics['rent_base'], "-", "Deducciones estatales y autonómicas"),
        ("Casilla 0465: Aportaciones a Planes de Pensiones", metrics['pension_base'], f"Reducción base: {metrics['pension_capped_base']}€", "Límite máximo de reducción 1.500€/año"),
        ("Casillas 0029-0036: Rendimientos del Capital Mobiliario", metrics['capital_income'], "-", "Dividendos e intereses bancarios"),
        ("COBERTURA DOCUMENTAL (JUSTIFICANTES Y FACTURAS)", f"{metrics['total_with_document']} de {metrics['total_transactions']} transacciones", f"{metrics['document_coverage_percentage']}% con factura", "Facturas guardadas con enlace a Gmail"),
    ]

    cur_row = 5
    for title, base, ded, note in rows_data:
        ws_summary.cell(row=cur_row, column=1, value=title).font = bold_font if "COBERTURA" in title else regular_font
        ws_summary.cell(row=cur_row, column=2, value=base).font = regular_font
        ws_summary.cell(row=cur_row, column=3, value=ded).font = regular_font
        ws_summary.cell(row=cur_row, column=4, value=note).font = regular_font

        for col_idx in range(1, 5):
            cell = ws_summary.cell(row=cur_row, column=col_idx)
            cell.border = thin_border
            if col_idx in (2, 3) and isinstance(cell.value, (int, float, Decimal)):
                cell.number_format = '#,##0.00 €'
                cell.alignment = Alignment(horizontal="right")
        cur_row += 1

    # Desglose por Categorías
    cur_row += 2
    ws_summary.cell(row=cur_row, column=1, value="DESGLOSE DETALLADO POR CATEGORÍAS").font = Font(bold=True, size=12, color="1E3A8A")
    cur_row += 1

    headers_cat = ["Categoría", "Casilla AEAT", "Tipo", "¿Deducible?", "Importe Total (€)", "Total Movimientos", "Con Factura"]
    for col_idx, h_text in enumerate(headers_cat, 1):
        c = ws_summary.cell(row=cur_row, column=col_idx, value=h_text)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center" if col_idx in (3, 4, 6, 7) else ("right" if col_idx == 5 else "left"))

    cur_row += 1
    for cat in summary['categories_breakdown']:
        ws_summary.cell(row=cur_row, column=1, value=cat['category_name']).font = regular_font
        ws_summary.cell(row=cur_row, column=2, value=cat['aeat_code'] or '-').font = regular_font
        ws_summary.cell(row=cur_row, column=3, value='Ingreso' if cat['is_income'] else 'Gasto').font = regular_font
        ws_summary.cell(row=cur_row, column=4, value='Sí' if cat['is_deductible'] else 'No').font = regular_font
        c_amt = ws_summary.cell(row=cur_row, column=5, value=float(cat['total_amount']))
        c_amt.font = regular_font
        c_amt.number_format = '#,##0.00 €'
        ws_summary.cell(row=cur_row, column=6, value=cat['count']).font = regular_font
        ws_summary.cell(row=cur_row, column=7, value=f"{cat['with_document_count']} / {cat['count']}").font = regular_font

        for col_idx in range(1, 8):
            ws_summary.cell(row=cur_row, column=col_idx).border = thin_border
        cur_row += 1

    # Autoajuste de anchos de columna
    for col in ws_summary.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # ----------------------------------------------------
    # HOJA 2: Movimientos Detallados con Enlaces a Facturas
    # ----------------------------------------------------
    ws_details = wb.create_sheet(title="Movimientos Detallados")
    ws_details.views.sheetView[0].showGridLines = True

    txs = Transaction.objects.filter(
        user=user,
        is_deleted=False
    ).filter(
        Q(tax_year=year) | (Q(tax_year__isnull=True) & Q(date__year=year))
    ).select_related('category', 'account').prefetch_related('documents').order_by('date')

    tx_headers = [
        "Fecha", "Concepto / Descripción", "Importe (€)", "Tipo", "Categoría",
        "Casilla AEAT", "¿Deducible IRPF?", "¿Tiene Factura?", "Nombre Documento", "Enlace Gmail"
    ]

    for col_idx, h_text in enumerate(tx_headers, 1):
        c = ws_details.cell(row=1, column=col_idx, value=h_text)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center" if col_idx in (1, 4, 7, 8) else ("right" if col_idx == 3 else "left"))

    detail_row = 2
    for tx in txs:
        doc = tx.documents.first()
        is_deductible = tx.is_tax_deductible or (tx.category and tx.category.tax_deductible)

        ws_details.cell(row=detail_row, column=1, value=tx.date.strftime('%d/%m/%Y')).alignment = Alignment(horizontal="center")
        ws_details.cell(row=detail_row, column=2, value=tx.description)
        
        amt_cell = ws_details.cell(row=detail_row, column=3, value=float(tx.amount))
        amt_cell.number_format = '#,##0.00 €'
        amt_cell.alignment = Alignment(horizontal="right")

        ws_details.cell(row=detail_row, column=4, value='Ingreso' if tx.type == 'income' else 'Gasto').alignment = Alignment(horizontal="center")
        ws_details.cell(row=detail_row, column=5, value=tx.category.name if tx.category else 'Sin categoría')
        ws_details.cell(row=detail_row, column=6, value=tx.category.aeat_code if tx.category and tx.category.aeat_code else '-')
        ws_details.cell(row=detail_row, column=7, value='Sí' if is_deductible else 'No').alignment = Alignment(horizontal="center")
        ws_details.cell(row=detail_row, column=8, value='Sí' if doc else 'No').alignment = Alignment(horizontal="center")
        ws_details.cell(row=detail_row, column=9, value=doc.file_name if doc else '-')
        
        gmail_cell = ws_details.cell(row=detail_row, column=10, value=doc.gmail_web_link if (doc and doc.gmail_web_link) else '-')
        if doc and doc.gmail_web_link:
            gmail_cell.hyperlink = doc.gmail_web_link
            gmail_cell.font = Font(color="2563EB", underline="single")

        for col_idx in range(1, 11):
            ws_details.cell(row=detail_row, column=col_idx).border = thin_border
            if col_idx != 10:
                ws_details.cell(row=detail_row, column=col_idx).font = regular_font
        detail_row += 1

    for col in ws_details.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_details.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def export_tax_csv(user, year: int) -> io.StringIO:
    """Genera exportación en formato CSV limpio."""
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')
    writer.writerow([
        "Fecha", "Concepto", "Importe", "Tipo", "Categoria", "Casilla_AEAT",
        "Deducible_IRPF", "Tiene_Factura", "Nombre_Documento", "Enlace_Gmail"
    ])

    txs = Transaction.objects.filter(
        user=user,
        is_deleted=False
    ).filter(
        Q(tax_year=year) | (Q(tax_year__isnull=True) & Q(date__year=year))
    ).select_related('category').prefetch_related('documents').order_by('date')

    for tx in txs:
        doc = tx.documents.first()
        is_deductible = tx.is_tax_deductible or (tx.category and tx.category.tax_deductible)
        writer.writerow([
            tx.date.isoformat(),
            tx.description,
            str(tx.amount),
            tx.type,
            tx.category.name if tx.category else '',
            tx.category.aeat_code if tx.category and tx.category.aeat_code else '',
            'SI' if is_deductible else 'NO',
            'SI' if doc else 'NO',
            doc.file_name if doc else '',
            doc.gmail_web_link if (doc and doc.gmail_web_link) else ''
        ])

    output.seek(0)
    return output

